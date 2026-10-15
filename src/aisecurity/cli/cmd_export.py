# -*- coding: utf-8 -*-
"""
aisec export -- export case reports.
Supported formats:
    markdown  -- Markdown table
    json      -- structured JSON
    docx      -- Word (requires python-docx)
    xlsx      -- Excel (requires openpyxl)
    html      -- HTML
"""
import json
from ..core.registry import CATEGORY_REGISTRY

def add_export_parser(sub):
    p = sub.add_parser("export", help="export case reports")
    p.add_argument("--format", "-f", default="markdown",
                   choices=["markdown", "json", "docx", "xlsx", "html"],
                   help="output format")
    p.add_argument("--output", "-o", default=None,
                   help="output file; if not set, prints to stdout "
                        "(docx/xlsx require this to be set)")


# ---------- Markdown ----------
def _escape_cell(s):
    if s is None:
        return ""
    s = str(s).replace("|", "\\|").replace("\n", " ").replace("\r", "")
    return s.strip()


def _to_markdown():
    lines = ["# AISecurityTest Case Library", ""]
    for name, info in CATEGORY_REGISTRY.items():
        lines.append("## %s %s" % (info["icon"], name))
        lines.append("")
        lines.append("> %s" % info.get("desc", ""))
        lines.append("")
        lines.append("| Case ID | Subcategory (L3) | Level | OWASP | Note |")
        lines.append("|---|---|---|---|---|")
        for cid, case in info["cases"].items():
            cells = [
                _escape_cell(cid),
                _escape_cell(case.get("sub3", "")),
                _escape_cell(case.get("level", "")),
                _escape_cell(case.get("owasp", "")),
                _escape_cell(case.get("note", "")),
            ]
            lines.append("| %s |" % " | ".join(cells))
        lines.append("")
    return "\n".join(lines)


# ---------- JSON ----------
def _to_json():
    data = {}
    for name, info in CATEGORY_REGISTRY.items():
        data[name] = {
            "icon": info["icon"],
            "desc": info.get("desc", ""),
            "order": info["order"],
            "cases": info["cases"],
        }
    return json.dumps(data, ensure_ascii=False, indent=2)


# ---------- Collect all records ----------
def _all_records():
    records = []
    for name, info in CATEGORY_REGISTRY.items():
        for cid, case in info["cases"].items():
            records.append({
                "case_id": cid,
                "cat": case.get("cat", name),
                "sub": case.get("sub", ""),
                "sub3": case.get("sub3", ""),
                "pre": case.get("pre", ""),
                "steps": case.get("steps", []),
                "expect": case.get("expect", ""),
                "level": case.get("level", ""),
                "defense": case.get("defense", ""),
                "owasp": case.get("owasp", ""),
                "note": case.get("note", ""),
            })
    return records


# ---------- docx ----------
def _to_docx(records, path):
    try:
        from docx import Document
    except ImportError:
        raise RuntimeError("python-docx not installed. Run: pip install python-docx")
    doc = Document()
    doc.add_heading("AISecurityTest - Case Library", 0)
    doc.add_paragraph("Total cases: %d" % len(records))
    for r in records:
        doc.add_heading("%s  %s" % (r["case_id"], r["sub3"]), level=1)
        doc.add_paragraph("")
        doc.add_paragraph("Category (L1): %s" % r["cat"])
        doc.add_paragraph("Category (L2): %s" % r["sub"])
        doc.add_paragraph("Level: %s" % r["level"])
        doc.add_paragraph("OWASP: %s" % r["owasp"])
        doc.add_paragraph("Preconditions: %s" % r["pre"])
        doc.add_paragraph("Test steps:")
        for i, s in enumerate(r["steps"], 1):
            doc.add_paragraph("%d. %s" % (i, s), style="List Bullet")
        doc.add_paragraph("Expected: %s" % r["expect"])
        doc.add_paragraph("Defense: %s" % r["defense"])
        doc.add_paragraph("Note: %s" % r["note"])
    doc.save(path)
    return path


# ---------- xlsx ----------
HEADERS = ["Case ID", "Category (L1)", "Category (L2)", "Category (L3)",
           "Preconditions", "Test Steps", "Expected", "Level",
           "Defense", "OWASP", "Note"]


def _to_xlsx(records, path):
    try:
        import openpyxl
        from openpyxl.styles import Font, Alignment, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError:
        raise RuntimeError("openpyxl not installed. Run: pip install openpyxl")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AI Security Test Cases"
    ws.append(HEADERS)
    for c in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for r in records:
        steps = "\n".join("%d. %s" % (i, s) for i, s in enumerate(r["steps"], 1))
        ws.append([r["case_id"], r["cat"], r["sub"], r["sub3"],
                   r["pre"], steps, r["expect"], r["level"],
                   r["defense"], r["owasp"], r["note"]])
    for i, w in enumerate([10, 16, 12, 24, 36, 44, 36, 10, 44, 14, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A2"
    wb.save(path)
    return path


# ---------- html ----------
def _to_html(records):
    rows = ""
    for r in records:
        steps = "<br>".join("%d. %s" % (i, s) for i, s in enumerate(r["steps"], 1))
        rows += ("<tr><td>%s</td><td>%s</td><td>%s</td><td>%s</td>"
                 "<td>%s</td><td>%s</td><td>%s</td><td>%s</td>"
                 "<td>%s</td><td>%s</td><td>%s</td></tr>") % (
            r["case_id"], r["cat"], r["sub"], r["sub3"],
            r["pre"], steps, r["expect"], r["level"],
            r["defense"], r["owasp"], r["note"])

    # Use .format and double braces to escape CSS.
    template = """<html><head><meta charset="utf-8"><title>AISecurityTest Case Library</title>
<style>
body {{ font-family: Segoe UI, Arial, sans-serif; padding: 20px; }}
table {{ border-collapse: collapse; width: 100%; }}
th, td {{ border: 1px solid #999; padding: 6px; vertical-align: top; font-size: 12px; }}
th {{ background: #4472C4; color: #fff; }}
</style></head><body>
<h1>AISecurityTest Case Library</h1>
<p>Total cases: {count}</p>
<table><tr>{headers}</tr>{rows}</table></body></html>"""

    return template.format(
        count=len(records),
        headers="".join("<th>%s</th>" % h for h in HEADERS),
        rows=rows,
    )


def handle_export(args):
    fmt = args.format

    if fmt == "markdown":
        text = _to_markdown()
    elif fmt == "json":
        text = _to_json()
    elif fmt == "html":
        text = _to_html(_all_records())
    elif fmt in ("docx", "xlsx"):
        if not args.output:
            print("--format %s requires --output file path." % fmt)
            return 1
        records = _all_records()
        if fmt == "docx":
            _to_docx(records, args.output)
        else:
            _to_xlsx(records, args.output)
        print("Exported: %s (%d cases)" % (args.output, len(records)))
        return 0
    else:
        print("Unknown format: %s" % fmt)
        return 1

    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(text)
        print("Exported: %s" % args.output)
    else:
        print(text)
    return 0